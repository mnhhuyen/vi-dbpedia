from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Bao_cao.md"
OUTPUT = ROOT / "Bao_cao_VI_DBPEDIA.docx"
NAVY = RGBColor(31, 78, 121)


def set_cell(cell, value, header=False):
    cell.text = value.replace("`", "")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            if header:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    if header:
        shading = "1F4E79"
        cell._tc.get_or_add_tcPr().append(
            __import__("docx").oxml.parse_xml(
                f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{shading}"/>'
            )
        )


def add_text(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part.strip("`") if part.startswith("`") else part)


def add_table(document, rows):
    cells = [[value.strip() for value in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r"[-: ]+", value) for value in row)]
    if not cells:
        return
    table = document.add_table(rows=1, cols=len(cells[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, value in enumerate(cells[0]):
        set_cell(table.rows[0].cells[index], value, True)
    for row in cells[1:]:
        target = table.add_row().cells
        for index, value in enumerate(row[:len(target)]):
            set_cell(target[index], value)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.keep_with_next = True

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title_done = False
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
        if image_match:
            alt, image_path = image_match.groups()
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(str(ROOT / image_path), width=Cm(15.5))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(alt)
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(9)
            index += 1
            continue
        if line.startswith("| "):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue
        if line.startswith("# "):
            if title_done:
                document.add_page_break()
            paragraph = document.add_paragraph(style="Title" if not title_done else "Heading 1")
            add_text(paragraph, line[2:])
            if not title_done:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.runs[0].font.name = "Times New Roman"
                paragraph.runs[0].font.size = Pt(22)
                paragraph.runs[0].font.bold = True
                paragraph.runs[0].font.color.rgb = NAVY
                title_done = True
            index += 1
            continue
        if line.startswith("## "):
            if line.startswith("## Tóm tắt"):
                document.add_page_break()
            paragraph = document.add_paragraph(style="Heading 1")
            add_text(paragraph, line[3:])
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_text(paragraph, line[4:])
            index += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            add_text(paragraph, line[2:])
            index += 1
            continue
        if re.match(r"\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            add_text(paragraph, re.sub(r"^\d+\. ", "", line))
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_text(paragraph, line)
        index += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("VI-DBpedia · Báo cáo môn học     |     ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    field = __import__("docx").oxml.parse_xml(
        '<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="PAGE"/>'
    )
    footer._p.append(field)
    document.core_properties.title = "Xây dựng phiên bản DBpedia cho tiếng Việt"
    document.core_properties.subject = "Báo cáo môn học Semantic Web"
    document.core_properties.author = "Sinh viên"
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
